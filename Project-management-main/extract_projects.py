import json
from datetime import datetime
from openpyxl import load_workbook

SOURCE = "project_sources/01-iLD-Project-tracking-2026.xlsx"
OUTPUT = "site/data.js"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def clean(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        return ""
    return value


wb = load_workbook(SOURCE, data_only=False)
ws = wb["Factory Project "]
projects = []
for r in range(3, ws.max_row + 1):
    number = clean(ws.cell(r, 4).value)
    name = clean(ws.cell(r, 5).value)
    if number == "" and name == "":
        continue
    monthly = {MONTHS[i]: clean(ws.cell(r, 21 + i).value) for i in range(12)}
    projects.append({
        "id": str(number or f"row-{r}"),
        "keyDriver": clean(ws.cell(r, 1).value),
        "factoryTarget": clean(ws.cell(r, 2).value),
        "targetValue": clean(ws.cell(r, 3).value),
        "number": number,
        "name": name,
        "keyActivity": clean(ws.cell(r, 6).value),
        "leader": clean(ws.cell(r, 7).value),
        "department": clean(ws.cell(r, 8).value),
        "supporter": clean(ws.cell(r, 9).value),
        "teamMember": clean(ws.cell(r, 10).value),
        "costSaving": clean(ws.cell(r, 11).value),
        "costInvestment": clean(ws.cell(r, 12).value),
        "charter": clean(ws.cell(r, 13).value),
        "charterCompleted": clean(ws.cell(r, 14).value),
        "target2026": clean(ws.cell(r, 15).value),
        "timing": clean(ws.cell(r, 16).value),
        "status": clean(ws.cell(r, 17).value),
        "coach": clean(ws.cell(r, 18).value),
        "sponsor": clean(ws.cell(r, 19).value),
        "priority": clean(ws.cell(r, 20).value),
        "monthly": monthly,
        "q1InitialVolume": clean(ws.cell(r, 33).value),
        "q1ActualVolume": clean(ws.cell(r, 34).value),
        "q1InitialCost": clean(ws.cell(r, 35).value),
        "q1ActualCost": clean(ws.cell(r, 36).value),
        "q2InitialVolume": clean(ws.cell(r, 37).value),
        "q2ActualVolume": clean(ws.cell(r, 38).value),
        "q2InitialCost": clean(ws.cell(r, 39).value),
        "q2ActualCost": clean(ws.cell(r, 40).value),
        "q3InitialVolume": clean(ws.cell(r, 41).value),
        "q3ActualVolume": clean(ws.cell(r, 42).value),
        "q3InitialCost": clean(ws.cell(r, 43).value),
        "q3ActualCost": clean(ws.cell(r, 44).value),
        "q4InitialVolume": clean(ws.cell(r, 45).value),
        "q4ActualVolume": clean(ws.cell(r, 46).value),
        "q4InitialCost": clean(ws.cell(r, 47).value),
        "q4ActualCost": clean(ws.cell(r, 48).value),
        "fyInitialVolume": clean(ws.cell(r, 49).value),
        "fyActualVolume": clean(ws.cell(r, 50).value),
        "fyInitialCost": clean(ws.cell(r, 51).value),
        "fyActualCost": clean(ws.cell(r, 52).value),
    })

mom_ws = wb["MoM"]
meetings = []
for r in range(3, mom_ws.max_row + 1):
    if mom_ws.cell(r, 1).value is None:
        continue
    meetings.append({
        "id": str(clean(mom_ws.cell(r, 1).value)),
        "reviewDate": clean(mom_ws.cell(r, 2).value),
        "project": clean(mom_ws.cell(r, 3).value),
        "activities": clean(mom_ws.cell(r, 4).value),
        "leader": clean(mom_ws.cell(r, 5).value),
        "comment": clean(mom_ws.cell(r, 6).value),
        "timing": clean(mom_ws.cell(r, 7).value),
        "status": clean(mom_ws.cell(r, 8).value),
        "finishTime": clean(mom_ws.cell(r, 9).value),
    })

payload = {"source": "iLD Project tracking 2026.xlsx", "projects": projects, "meetings": meetings}
with open(OUTPUT, "w", encoding="utf-8") as f:
    f.write("window.ILD_DATA = ")
    json.dump(payload, f, ensure_ascii=False, indent=2)
    f.write(";\n")
print(f"Extracted {len(projects)} projects and {len(meetings)} meeting records")
